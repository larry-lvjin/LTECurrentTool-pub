import openpyxl
import sys

def analyze_excel(filename):
    wb = openpyxl.load_workbook(filename)
    ws = wb.active

    print(f"\n{'='*80}")
    print(f"文件: {filename}")
    print(f"{'='*80}")

    # 读取表头
    headers = []
    for col in range(1, ws.max_column + 1):
        cell_value = ws.cell(1, col).value
        if cell_value:
            headers.append(str(cell_value))

    print(f"列数: {len(headers)}")
    print(f"行数: {ws.max_row}")
    print(f"\n表头: {', '.join(headers[:10])}")  # 显示前10列

    # 读取前5行数据示例
    print(f"\n前5行数据示例:")
    for row in range(2, min(7, ws.max_row + 1)):
        row_data = []
        for col in range(1, min(11, ws.max_column + 1)):  # 前10列
            cell_value = ws.cell(row, col).value
            row_data.append(str(cell_value) if cell_value is not None else "")
        print(f"  行{row}: {' | '.join(row_data)}")

    # 找到关键列索引
    avg_col = None
    max_col = None
    min_col = None
    samples_col = None
    band_col = None
    channel_col = None
    tpc_col = None

    for i, header in enumerate(headers, 1):
        if '平均' in header or 'mA' in header and i <= 10:
            if avg_col is None:
                avg_col = i
        if '最大' in header:
            max_col = i
        if '最小' in header:
            min_col = i
        if '样本数' in header:
            samples_col = i
        if 'BAND' in header:
            band_col = i
        if '信道' in header and '类型' not in header:
            channel_col = i
        if '测试类型' in header:
            tpc_col = i

    print(f"\n关键列位置:")
    print(f"  BAND列: {band_col}")
    print(f"  信道列: {channel_col}")
    print(f"  测试类型列: {tpc_col}")
    print(f"  平均电流列: {avg_col}")
    print(f"  最大电流列: {max_col}")
    print(f"  最小电流列: {min_col}")
    print(f"  样本数列: {samples_col}")

    # 统计数据
    data_rows = []
    for row in range(2, ws.max_row + 1):
        band = ws.cell(row, band_col).value if band_col else None
        channel = ws.cell(row, channel_col).value if channel_col else None
        tpc = ws.cell(row, tpc_col).value if tpc_col else None
        avg = ws.cell(row, avg_col).value if avg_col else None
        max_val = ws.cell(row, max_col).value if max_col else None
        min_val = ws.cell(row, min_col).value if min_col else None
        samples = ws.cell(row, samples_col).value if samples_col else None

        if avg is not None:
            # 解析样本数（可能是"1024/1024"格式）
            sample_count = 0
            if samples:
                if isinstance(samples, str) and '/' in str(samples):
                    sample_count = int(str(samples).split('/')[0])
                elif isinstance(samples, (int, float)):
                    sample_count = int(samples)
                else:
                    try:
                        sample_count = int(samples)
                    except:
                        sample_count = 0

            data_rows.append({
                'band': band,
                'channel': channel,
                'tpc': tpc,
                'avg': float(avg) if avg else 0,
                'max': float(max_val) if max_val else 0,
                'min': float(min_val) if min_val else 0,
                'samples': sample_count
            })

    print(f"\n数据统计:")
    print(f"  有效数据行数: {len(data_rows)}")

    if data_rows:
        # 样本数统计
        sample_counts = [row['samples'] for row in data_rows if row['samples'] > 0]
        if sample_counts:
            print(f"  样本数范围: {min(sample_counts)} ~ {max(sample_counts)}")
            print(f"  平均样本数: {sum(sample_counts) / len(sample_counts):.1f}")

        # 平均电流统计
        avg_currents = [row['avg'] for row in data_rows if row['avg'] > 0]
        if avg_currents:
            print(f"  平均电流范围: {min(avg_currents):.2f} ~ {max(avg_currents):.2f} mA")

        # 检查异常大的最大值
        abnormal_max = [row for row in data_rows if row['max'] > 3000]
        if abnormal_max:
            print(f"\n  ⚠️ 发现异常大的最大值: {len(abnormal_max)}个")
            for row in abnormal_max[:5]:
                print(f"    {row['band']} {row['channel']} {row['tpc']}: Max={row['max']:.0f}mA, Avg={row['avg']:.2f}mA, Samples={row['samples']}")

    return data_rows

# 分析两个文件
data1 = analyze_excel(r"C:\Users\<user>\Desktop\222.xlsx")
data2 = analyze_excel(r"C:\Users\<user>\Desktop\333.xlsx")

# 对比分析
print(f"\n{'='*80}")
print(f"两次测试对比分析")
print(f"{'='*80}")

if data1 and data2:
    # 找到相同的测试点
    print(f"\n文件1数据点数: {len(data1)}")
    print(f"文件2数据点数: {len(data2)}")

    # 构建key用于匹配
    def make_key(row):
        return f"{row['band']}_{row['channel']}_{row['tpc']}"

    data1_dict = {make_key(row): row for row in data1}
    data2_dict = {make_key(row): row for row in data2}

    common_keys = set(data1_dict.keys()) & set(data2_dict.keys())
    print(f"共同测试点数: {len(common_keys)}")

    # 计算差异
    differences = []
    for key in common_keys:
        row1 = data1_dict[key]
        row2 = data2_dict[key]

        diff_pct = abs(row1['avg'] - row2['avg']) / ((row1['avg'] + row2['avg']) / 2) * 100 if row1['avg'] > 0 and row2['avg'] > 0 else 0

        differences.append({
            'key': key,
            'band': row1['band'],
            'channel': row1['channel'],
            'tpc': row1['tpc'],
            'avg1': row1['avg'],
            'avg2': row2['avg'],
            'diff': row2['avg'] - row1['avg'],
            'diff_pct': diff_pct,
            'samples1': row1['samples'],
            'samples2': row2['samples']
        })

    if differences:
        # 排序：差异百分比从大到小
        differences.sort(key=lambda x: x['diff_pct'], reverse=True)

        print(f"\n差异最大的10个测试点:")
        print(f"{'序号':<4} {'BAND':<6} {'信道':<8} {'TPC类型':<15} {'测试1(mA)':<12} {'测试2(mA)':<12} {'差异(mA)':<12} {'差异%':<10} {'样本1':<8} {'样本2':<8}")
        print("-" * 120)

        for i, diff in enumerate(differences[:10], 1):
            print(f"{i:<4} {str(diff['band']):<6} {str(diff['channel']):<8} {str(diff['tpc']):<15} "
                  f"{diff['avg1']:<12.2f} {diff['avg2']:<12.2f} {diff['diff']:<12.2f} "
                  f"{diff['diff_pct']:<10.1f} {diff['samples1']:<8} {diff['samples2']:<8}")

        # 统计差异分布
        diff_pcts = [d['diff_pct'] for d in differences]
        print(f"\n差异统计:")
        print(f"  平均差异: {sum(diff_pcts) / len(diff_pcts):.2f}%")
        print(f"  最大差异: {max(diff_pcts):.2f}%")
        print(f"  最小差异: {min(diff_pcts):.2f}%")

        # 差异分布
        large_diff = [d for d in differences if d['diff_pct'] > 10]
        medium_diff = [d for d in differences if 5 <= d['diff_pct'] <= 10]
        small_diff = [d for d in differences if d['diff_pct'] < 5]

        print(f"\n差异分布:")
        print(f"  差异>10%: {len(large_diff)}个 ({len(large_diff)/len(differences)*100:.1f}%)")
        print(f"  差异5-10%: {len(medium_diff)}个 ({len(medium_diff)/len(differences)*100:.1f}%)")
        print(f"  差异<5%: {len(small_diff)}个 ({len(small_diff)/len(differences)*100:.1f}%)")
