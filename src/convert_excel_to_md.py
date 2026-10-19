#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
将LTE电流规则.xlsx转换为Markdown格式
"""

from openpyxl import load_workbook
import sys

def excel_to_markdown(excel_path, md_path):
    """将Excel文件转换为Markdown格式"""

    # 读取Excel文件
    wb = load_workbook(excel_path, data_only=True)
    ws = wb.active  # 获取活动工作表

    # 创建Markdown内容
    md_content = []
    md_content.append("# LTE电流规则\n")
    md_content.append("> **注意**: 本文档为只读参考文档，请勿修改。\n")
    md_content.append("> 用于C# WPF控制CMW500和66319D电源的电流自动读取工具开发。\n\n")
    md_content.append("---\n\n")

    # 获取所有行数据
    rows = list(ws.iter_rows(values_only=True))

    if not rows:
        print("警告: Excel文件为空")
        return

    # 查找表格数据（跳过空行）
    table_data = []
    for row in rows:
        # 过滤掉完全为空的行
        if any(cell is not None and str(cell).strip() for cell in row):
            table_data.append(row)

    if not table_data:
        print("警告: 没有找到有效数据")
        return

    # 检查是否是表格格式（有标题行）
    # 假设第一行可能是标题
    max_cols = max(len(row) for row in table_data)

    # 生成Markdown表格
    for i, row in enumerate(table_data):
        # 补齐列数
        row_list = list(row) + [None] * (max_cols - len(row))

        # 转换单元格值为字符串
        cells = []
        for cell in row_list:
            if cell is None or str(cell).strip() == '':
                cells.append('')
            else:
                # 处理多行文本
                cell_str = str(cell).replace('\n', '<br>')
                cells.append(cell_str)

        # 生成表格行
        md_content.append('| ' + ' | '.join(cells) + ' |')

        # 如果是第一行，添加分隔线
        if i == 0:
            md_content.append('| ' + ' | '.join(['---'] * max_cols) + ' |')

    # 写入Markdown文件
    with open(md_path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(md_content))

    print(f"转换成功: {md_path}")
    print(f"共处理 {len(table_data)} 行数据")

if __name__ == '__main__':
    excel_file = r'D:\MyFile\CODE\VSCODE\LTECurrentTool\LTE电流规则.xlsx'
    md_file = r'D:\MyFile\CODE\VSCODE\LTECurrentTool\LTE电流规则.md'

    excel_to_markdown(excel_file, md_file)
